"""Vercel Serverless Function: slim PDF -> Word conversion.

Engine: PyMuPDF (text and image blocks) -> python-docx.
Why not pdf2docx here: its dependency bundle measures ~300 MB locally
(PyMuPDF 116 MB, opencv 72 MB, numpy 42 + 27 MB, fontTools 29 MB, ...),
over Vercel's 250 MB uncompressed function-package limit, and opencv
cannot be dropped because pdf2docx imports cv2 at the top level of
common/algorithm.py. This slim bundle (PyMuPDF + python-docx +
python-multipart) is roughly 130 MB and fits.

Output is text-focused: paragraphs keep font size / bold / italic from
the PDF spans and images are embedded at native size, but there is no
pdf2docx-style table reconstruction, so complex multi-column layouts
come out linearized.

Vercel Python runtime: the handler class must be named `handler` and
subclass http.server.BaseHTTPRequestHandler.
"""
import io
from http.server import BaseHTTPRequestHandler

MAX_BYTES = 50 * 1024 * 1024


def _is_cjk(ch):
    return "\u2e80" <= ch <= "\u9fff" or "\u3040" <= ch <= "\u30ff" or "\uac00" <= ch <= "\ud7af"


def convert_pdf_to_docx(pdf_bytes):
    import fitz
    from docx import Document
    from docx.shared import Pt

    src = fitz.open(stream=pdf_bytes, filetype="pdf")
    out = Document()
    npages = len(src)
    for pno, page in enumerate(src):
        d = page.get_text("dict")
        for b in d["blocks"]:
            if b["type"] == 1:
                # image block: this PyMuPDF version embeds the bytes
                # directly ("image"); older versions only give "xref".
                try:
                    data = b.get("image")
                    if not data and b.get("xref"):
                        data = src.extract_image(b["xref"])["image"]
                    if data:
                        p = out.add_paragraph()
                        p.add_run().add_picture(io.BytesIO(data))
                except Exception:
                    pass
                continue
            segments = []  # (text, size, bold, italic)
            lines = b.get("lines", [])
            for li, line in enumerate(lines):
                spans = line.get("spans", [])
                for span in spans:
                    text = span.get("text", "")
                    if text:
                        segments.append((
                            text,
                            span.get("size"),
                            bool(span.get("flags", 0) & 16),
                            bool(span.get("flags", 0) & 2),
                        ))
                # line separator: a space only between two non-CJK neighbours
                if li < len(lines) - 1 and segments:
                    cur_last = segments[-1][0][-1:]
                    nxt_spans = lines[li + 1].get("spans", [])
                    nxt_first = nxt_spans[0].get("text", "")[:1] if nxt_spans else ""
                    if cur_last and nxt_first and not _is_cjk(cur_last) and not _is_cjk(nxt_first):
                        segments.append((" ", None, False, False))
            if not segments:
                continue
            p = out.add_paragraph()
            for text, size, bold, italic in segments:
                run = p.add_run(text)
                if size:
                    try:
                        run.font.size = Pt(size)
                    except Exception:
                        pass
                run.bold = bold
                run.italic = italic
        if pno < npages - 1:
            out.add_page_break()
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def _read_upload(h):
    """Return (filename, pdf_bytes, error). error in {size, type, nofile}."""
    try:
        length = int(h.headers.get("Content-Length") or 0)
    except ValueError:
        length = 0
    if length <= 0 or length > MAX_BYTES:
        return None, None, "size"
    body = h.rfile.read(length)
    ctype = h.headers.get("Content-Type", "")
    if "multipart/form-data" in ctype:
        from multipart import parse_form

        got = {}

        def on_file(f):
            got["name"] = f.file_name
            got["obj"] = f.file_object

        parse_form({"Content-Type": ctype}, io.BytesIO(body), None, on_file)
        if "obj" not in got:
            return None, None, "nofile"
        got["obj"].seek(0)
        data = got["obj"].read()
        name = got["name"]
        if isinstance(name, bytes):
            name = name.decode("utf-8", "replace")
        return name, data, None
    if ctype.startswith("application/pdf"):
        name = h.headers.get("X-Filename", "document.pdf")
        return name, body, None
    return None, None, "type"


class handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def _send(self, code, body, ctype="application/json"):
        data = body if isinstance(body, bytes) else body.encode()
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        self._send(405, '{"error":"method not allowed"}')

    def do_POST(self):
        name, data, err = _read_upload(self)
        if err:
            code = 413 if err == "size" else 400
            self._send(code, '{"error":"invalid upload"}')
            return
        if not name or not name.lower().endswith(".pdf"):
            self._send(400, '{"error":"only PDF files are accepted"}')
            return
        if len(data) > MAX_BYTES:
            self._send(413, '{"error":"file exceeds 50 MB"}')
            return
        try:
            docx_bytes = convert_pdf_to_docx(data)
        except Exception as exc:  # noqa: BLE001 - surface as 500, no traceback
            self._send(500, '{"error":"conversion failed"}')
            return
        stem = name.rsplit(".", 1)[0] if "." in name else name
        safe = stem if stem.isascii() else "output"
        self.send_response(200)
        self.send_header(
            "Content-Type",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.send_header("Content-Disposition", 'attachment; filename="%s.docx"' % safe)
        self.send_header("Content-Length", str(len(docx_bytes)))
        self.end_headers()
        self.wfile.write(docx_bytes)
